import io
import logging
import math

import numpy as np
import pandas as pd
from flask import Flask, jsonify, render_template, request

from model.skill_analyzer import SkillGapAnalyzer

try:
    from pypdf import PdfReader
except ImportError:  # pragma: no cover
    PdfReader = None

try:
    from docx import Document
except ImportError:  # pragma: no cover
    Document = None

app = Flask(__name__)
logging.basicConfig(level=logging.INFO)

# Reject request bodies bigger than ~5 MB (resume uploads).
MAX_UPLOAD_BYTES = 5 * 1024 * 1024
app.config["MAX_CONTENT_LENGTH"] = MAX_UPLOAD_BYTES + 64 * 1024
ALLOWED_EXTENSIONS = {"pdf", "docx", "txt"}
MAX_PDF_PAGES = 20

# Initialize the Skill Gap Analyzer once at startup
analyzer = SkillGapAnalyzer()


# ---------------------------------------
# Helpers
# ---------------------------------------
def to_builtin(obj):
    """Convert numpy / pandas / NaN values into plain JSON-safe Python types."""
    if isinstance(obj, pd.DataFrame):
        return to_builtin(obj.to_dict(orient="records"))
    if isinstance(obj, pd.Series):
        return to_builtin(obj.to_dict())
    if isinstance(obj, dict):
        return {str(k): to_builtin(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple, set)):
        return [to_builtin(v) for v in obj]
    if isinstance(obj, np.ndarray):
        return to_builtin(obj.tolist())
    if isinstance(obj, np.integer):
        return int(obj)
    if isinstance(obj, (np.floating, float)):
        value = float(obj)
        return None if math.isnan(value) or math.isinf(value) else value
    if isinstance(obj, np.bool_):
        return bool(obj)
    return obj


def skills_to_string(value):
    """Accept either 'python, sql' or ['python', 'sql'] and return a clean string."""
    if isinstance(value, (list, tuple, set)):
        return ", ".join(str(s).strip() for s in value if str(s).strip())
    if isinstance(value, str):
        return value.strip()
    return ""


def error_response(message, status):
    return jsonify({"error": message}), status


@app.errorhandler(413)
def too_large(_err):
    return error_response("File is too large. Maximum size is 5 MB.", 413)


# ---------------------------------------
# Skill index (the analyzer owns the CSV now)
# ---------------------------------------
ALL_SKILLS = analyzer.get_known_skills()
ALL_SKILLS_LOWER = [s.lower() for s in ALL_SKILLS]


# ---------------------------------------
# Resume text extraction
# ---------------------------------------
class ResumeError(Exception):
    def __init__(self, message, status=400):
        super().__init__(message)
        self.message = message
        self.status = status


def extract_pdf_text(data):
    if PdfReader is None:
        raise ResumeError("PDF support is not installed. Run: pip install pypdf", 500)
    if not data.startswith(b"%PDF"):
        raise ResumeError("This file does not look like a valid PDF.")

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise ResumeError("This PDF is password protected. Please upload an unlocked copy.")

    parts = []
    for page in reader.pages[:MAX_PDF_PAGES]:
        parts.append(page.extract_text() or "")
    return "\n".join(parts)


def extract_docx_text(data):
    if Document is None:
        raise ResumeError("Word support is not installed. Run: pip install python-docx", 500)
    if not data.startswith(b"PK"):
        raise ResumeError("This file does not look like a valid .docx document.")

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def extract_resume_text(filename, data):
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in ALLOWED_EXTENSIONS:
        raise ResumeError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")

    try:
        if ext == "pdf":
            return extract_pdf_text(data)
        if ext == "docx":
            return extract_docx_text(data)
        return data.decode("utf-8", errors="ignore")
    except ResumeError:
        raise
    except Exception:
        app.logger.exception("Could not read resume file")
        raise ResumeError("Could not read this file. It may be corrupted.", 422)


# ---------------------------------------
# Home Page
# ---------------------------------------
@app.route("/")
def home():
    return render_template("index.html", roles=analyzer.get_roles())


# ---------------------------------------
# Skills: autocomplete suggestions
#   GET /api/skills/suggest?q=py&limit=8
# ---------------------------------------
@app.route("/api/skills/suggest", methods=["GET"])
def suggest_skills():
    q = request.args.get("q", "").strip().lower()
    try:
        limit = max(1, min(int(request.args.get("limit", 8)), 20))
    except ValueError:
        limit = 8

    if len(q) < 2:  # start suggesting after 2 characters
        return jsonify([])

    starts, contains = [], []
    for skill, low in zip(ALL_SKILLS, ALL_SKILLS_LOWER):
        if low.startswith(q):
            starts.append(skill)
        elif q in low:
            contains.append(skill)

    return jsonify((starts + contains)[:limit])


# ---------------------------------------
# Resume upload: extract skills from a file
#   POST /api/resume/extract   (multipart form, field name: "resume")
# ---------------------------------------
@app.route("/api/resume/extract", methods=["POST"])
def extract_resume():
    try:
        file = request.files.get("resume")
        if file is None or not file.filename:
            return error_response("Please choose a resume file to upload.", 400)

        data = file.read(MAX_UPLOAD_BYTES + 1)
        if not data:
            return error_response("The uploaded file is empty.", 400)
        if len(data) > MAX_UPLOAD_BYTES:
            return error_response("File is too large. Maximum size is 5 MB.", 413)

        # The file is processed in memory only and never written to disk.
        text = extract_resume_text(file.filename, data)

        if not text.strip():
            return error_response(
                "No readable text found. If this is a scanned or image-only "
                "resume, please upload a text-based PDF or a DOCX file.",
                422,
            )

        skills = analyzer.extract_skills(text)

        return jsonify({
            "filename": file.filename,
            "skills": skills,
            "count": len(skills),
        })

    except ResumeError as exc:
        return error_response(exc.message, exc.status)
    except Exception as exc:
        app.logger.exception("Error in /api/resume/extract")
        return error_response(str(exc), 500)


# ---------------------------------------
# Analyze Skill Gap
# ---------------------------------------
@app.route("/analyze", methods=["POST"])
def analyze():
    try:
        data = request.get_json(silent=True) or {}

        user_skills = skills_to_string(data.get("skills", ""))
        target_role = str(data.get("target_role", "")).strip()

        if not user_skills:
            return error_response("Please enter your skills.", 400)
        if not target_role:
            return error_response("Please select a target job role.", 400)

        result = analyzer.analyze(user_skills, target_role)

        if result is None:
            return error_response("Selected job role was not found.", 404)

        return jsonify(to_builtin(result))

    except Exception as exc:
        app.logger.exception("Error in /analyze")
        return error_response(str(exc), 500)


# ---------------------------------------
# Recommend Job Roles
# ---------------------------------------
@app.route("/recommend", methods=["POST"])
def recommend():
    try:
        data = request.get_json(silent=True) or {}

        user_skills = skills_to_string(data.get("skills", ""))
        if not user_skills:
            return error_response("Please enter your skills.", 400)

        results = analyzer.recommend_roles(user_skills)

        if isinstance(results, pd.DataFrame):
            results = results.copy()
            # similarity (0-1) -> percentage
            if "similarity" in results.columns:
                results["similarity"] = (results["similarity"] * 100).round(2)

        return jsonify(to_builtin(results))

    except Exception as exc:
        app.logger.exception("Error in /recommend")
        return error_response(str(exc), 500)


# ---------------------------------------
# API: Get Available Job Roles
# ---------------------------------------
@app.route("/api/roles", methods=["GET"])
def get_roles():
    return jsonify({"roles": analyzer.get_roles()})


# ---------------------------------------
# API: Skills needed for one role (Roles page detail view)
#   GET /api/roles/Python%20Developer
#   GET /api/roles/Python%20Developer?have=Python,Git
#
# With ?have=..., the response also splits the role's skills into
# "have" and "missing" (aliases like "js" count as "JavaScript").
# ---------------------------------------
@app.route("/api/roles/<path:role>", methods=["GET"])
def get_role_detail(role):
    info = analyzer.get_role_skills(role)
    if info is None:
        return error_response("Job role not found.", 404)

    skills = info["skills"]
    response = {"job_role": info["job_role"], "skills": skills}

    have_param = request.args.get("have", "").strip()
    if have_param:
        have_set = analyzer.clean_skills(have_param)
        have, missing = [], []
        for skill in skills:
            if analyzer.normalize_skill(skill) in have_set:
                have.append(skill)
            else:
                missing.append(skill)
        response["have"] = have
        response["missing"] = missing

    return jsonify(response)


# ---------------------------------------
# Run Flask Application
# ---------------------------------------
if __name__ == "__main__":
    app.run(debug=True, host="127.0.0.1", port=5000)